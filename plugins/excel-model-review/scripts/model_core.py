import csv, math, openpyxl
from collections import Counter
from openpyxl.utils import get_column_letter, range_boundaries

def load_book(p): return openpyxl.load_workbook(p,data_only=False,keep_vba=str(p).lower().endswith('.xlsm'))
def formula(v): return isinstance(v,str) and v.startswith('=')
def kind(a,b):
    if a is None:return 'added'
    if b is None:return 'removed'
    if formula(a) and formula(b):return 'formula_changed'
    if formula(a):return 'formula_to_value'
    if formula(b):return 'value_to_formula'
    return 'value_changed'
def changed(a,b,at=0,rt=0):
    out=[]
    for r in range(1,max(a.max_row,b.max_row)+1):
      for c in range(1,max(a.max_column,b.max_column)+1):
        x,y=a.cell(r,c).value,b.cell(r,c).value
        eq=math.isclose(float(x),float(y),abs_tol=at,rel_tol=rt) if type(x) in (int,float) and type(y) in (int,float) else x==y
        if not eq: out.append({'sheet':b.title,'row':r,'col':c,'cell':f'{get_column_letter(c)}{r}','change_type':kind(x,y),'before':x,'after':y})
    return out
def clusters(ch,gap_r=2,gap_c=2):
    rem=set(range(len(ch))); out=[]
    while rem:
      q=[rem.pop()]; group=[]
      while q:
        i=q.pop(); group.append(ch[i])
        near=[j for j in rem if abs(ch[i]['row']-ch[j]['row'])<=gap_r and abs(ch[i]['col']-ch[j]['col'])<=gap_c]
        for j in near: rem.remove(j); q.append(j)
      out.append(group)
    return out
def text(c):
    v=c.value
    return isinstance(v,str) and v.strip() and not v.startswith('=') and len(v.strip())<=120 and not(v.strip().isdigit() and len(v.strip())<=4)
def label(ws,r1,c1,r2,c2):
    cand=[]
    def add(r,c,score,src):
      if 1<=r<=ws.max_row and 1<=c<=ws.max_column and text(ws.cell(r,c)):
        x=ws.cell(r,c); score+=14 if x.font.bold else 0
        score+=10 if any(z.min_row<=r<=z.max_row and z.min_col<=c<=z.max_col for z in ws.merged_cells.ranges) else 0
        cand.append((score,str(x.value).strip(),src))
    for d in range(1,7):
      for r in range(r1,min(r2,r1+4)+1): add(r,c1-d,90-d*8,'left')
    for d in range(1,8):
      for c in range(max(1,c1-3),c2+1): add(r1-d,c,82-d*7,'above')
    for r in range(r1,min(r2,r1+3)+1):
      for c in range(c1,min(c2,c1+4)+1): add(r,c,72,'inside')
    if not cand:return '(sem rotulo identificado)','fallback'
    cand.sort(reverse=True); return cand[0][1],cand[0][2]
def area(g,old,new):
    r1,r2=min(x['row'] for x in g),max(x['row'] for x in g); c1,c2=min(x['col'] for x in g),max(x['col'] for x in g)
    lab,src=label(new,r1,c1,r2,c2)
    if lab.startswith('(sem'): lab,src=label(old,r1,c1,r2,c2)
    ct=Counter(x['change_type'] for x in g)
    cls='logica/hardcode' if ct['formula_to_value'] else 'logica de calculo' if ct['formula_changed'] or ct['value_to_formula'] else 'bloco incluido/expandido' if ct['added']>=2 else 'bloco removido/reduzido' if ct['removed']>=2 else 'conteudo/premissas'
    note='; '.join(f'{ct[k]} {k}' for k in ct)
    return {'sheet':new.title,'range':f'{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}','label':lab,'label_source':src,'change_class':cls,'severity':'critical' if ct['formula_to_value'] else 'warning' if ct['formula_changed'] or ct['removed'] else 'info','changed_cells':len(g),'formulas_changed':ct['formula_changed'],'formula_to_value':ct['formula_to_value'],'value_to_formula':ct['value_to_formula'],'values_changed':ct['value_changed'],'added':ct['added'],'removed':ct['removed'],'note':note,'evidence_cells':[x['cell'] for x in g[:30]]}
def compare_pair(oldp,newp,row_gap=2,col_gap=2,abs_tol=0,rel_tol=0):
    A,B=load_book(oldp),load_book(newp); areas=[]; evidence=[]
    old,new=set(A.sheetnames),set(B.sheetnames)
    for s in A.sheetnames:
      if s not in new: continue
      ch=changed(A[s],B[s],abs_tol,rel_tol); evidence+=ch; areas += [area(g,A[s],B[s]) for g in clusters(ch,row_gap,col_gap)]
    for s in sorted(new-old): areas.append({'sheet':s,'range':'A1','label':'nova aba','label_source':'new_sheet','change_class':'aba incluida','severity':'info','changed_cells':0,'formulas_changed':0,'formula_to_value':0,'value_to_formula':0,'values_changed':0,'added':0,'removed':0,'note':'worksheet incluida','evidence_cells':[]})
    for s in sorted(old-new): areas.append({'sheet':s,'range':'A1','label':'aba removida','label_source':'removed_sheet','change_class':'aba removida','severity':'warning','changed_cells':0,'formulas_changed':0,'formula_to_value':0,'value_to_formula':0,'values_changed':0,'added':0,'removed':0,'note':'worksheet removida','evidence_cells':[]})
    return {'metadata':{'old_file':str(oldp),'new_file':str(newp)},'summary':{'areas_changed':len(areas),'cell_evidence_count':len(evidence),'sheets_added':sorted(new-old),'sheets_removed':sorted(old-new),'sheet_order_changed':A.sheetnames!=B.sheetnames},'areas':areas,'cell_evidence':evidence}
def write_area_csv(r,p):
    f=['sheet','range','label','label_source','change_class','severity','changed_cells','formulas_changed','formula_to_value','value_to_formula','values_changed','added','removed','note']; w=csv.DictWriter(open(p,'w',newline='',encoding='utf-8-sig'),fieldnames=f); w.writeheader(); [w.writerow({k:a.get(k) for k in f}) for a in r['areas']]
def write_evidence_csv(r,p):
    f=['sheet','cell','change_type','before','after']; w=csv.DictWriter(open(p,'w',newline='',encoding='utf-8-sig'),fieldnames=f); w.writeheader(); [w.writerow({k:a.get(k) for k in f}) for a in r['cell_evidence']]
def write_pair_markdown(r,p):
    L=['# Excel Model Structural Diff','','| Aba | Area | Rotulo inferido | Tipo | Evidencia |','|---|---|---|---|---|']+[f"| {a['sheet']} | {a['range']} | {a['label']} | {a['change_class']} | {a['note']} |" for a in r['areas']]; p.write_text('\n'.join(L),encoding='utf-8')
def related(a,b):
    if a['sheet']!=b['sheet']:return False
    if a['label']==b['label'] and not a['label'].startswith('(sem'):return True
    ac1,ar1,ac2,ar2=range_boundaries(a['range']); bc1,br1,bc2,br2=range_boundaries(b['range']); return not(ar2<br1-2 or br2<ar1-2 or ac2<bc1-2 or bc2<ac1-2)
def consolidate_evolution(reports,names):
    T=[]
    for i,r in enumerate(reports):
      for a in r['areas']:
        t=next((x for x in T if related(x['rep'],a)),None); e={'transition':f'{names[i]} -> {names[i+1]}',**a}
        if t:t['events'].append(e)
        else:T.append({'sheet':a['sheet'],'label':a['label'],'rep':a,'events':[e]})
    for t in T:
      t['dominant_change']=Counter(e['change_class'] for e in t['events']).most_common(1)[0][0]; t['first_seen']=t['events'][0]['transition']; t['last_seen']=t['events'][-1]['transition']; t['event_count']=len(t['events']); t.pop('rep')
    return {'versions':names,'threads':T,'pair_reports':reports}
def write_evolution_markdown(r,p):
    L=['# Excel Model Evolution','','**Versoes:** '+' -> '.join(r['versions']),'']
    for t in r['threads']:
      L += [f"## {t['sheet']} — {t['label']}"]+[f"- {e['transition']} — {e['range']} — {e['change_class']}: {e['note']}" for e in t['events']]+['']
    p.write_text('\n'.join(L),encoding='utf-8')
