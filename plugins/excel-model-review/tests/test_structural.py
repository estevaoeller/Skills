from pathlib import Path
import json, subprocess, sys, shutil
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tests" / "_tmp"


def make_book(path: Path, version: int):
    wb = Workbook()
    ws = wb.active
    ws.title = "CAPEX"
    ws["B2"] = "Cronograma de Investimentos"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B5"] = "Obras civis"
    ws["C4"], ws["D4"], ws["E4"] = "2026", "2027", "2028"
    ws["C5"], ws["D5"], ws["E5"] = 100, 120, 130
    ws["B9"] = "Depreciacao"
    ws["B9"].font = Font(bold=True)
    ws["C10"], ws["D10"] = "=C5/10", "=D5/10"
    if version >= 2:
        ws["D5"], ws["E5"], ws["F4"], ws["F5"] = 150, 180, "2029", 200
        ws["D10"] = "=D5/8"
    if version >= 3:
        ws["C14"] = "Sensibilidade CAPEX"
        ws["C14"].font = Font(bold=True, size=12)
        ws["C15"], ws["D15"], ws["E15"] = 0.9, 1.0, 1.1
        ws["C16"], ws["D16"], ws["E16"] = "=SUM(C5:F5)*C15", "=SUM(C5:F5)*D15", "=SUM(C5:F5)*E15"
    wb.save(path)


def main():
    if TMP.exists(): shutil.rmtree(TMP)
    TMP.mkdir(parents=True)
    v1, v2, v3 = TMP/"v1.xlsx", TMP/"v2.xlsx", TMP/"v3.xlsx"
    make_book(v1, 1); make_book(v2, 2); make_book(v3, 3)
    diff_out = TMP/"diff"
    subprocess.check_call([sys.executable, str(ROOT/"scripts"/"compare_structural.py"), str(v1), str(v2), "--out", str(diff_out)])
    diff = json.loads((diff_out/"diff.json").read_text(encoding="utf-8"))
    assert diff["summary"]["areas_changed"] >= 2
    labels = [a["label"] for a in diff["areas"]]
    assert any("Obras civis" in x or "Cronograma de Investimentos" in x for x in labels), labels
    assert any("Depreciacao" in x for x in labels), labels
    evo_out = TMP/"evo"
    subprocess.check_call([sys.executable, str(ROOT/"scripts"/"evolution.py"), str(v1), str(v2), str(v3), "--out", str(evo_out)])
    evo = json.loads((evo_out/"evolution.json").read_text(encoding="utf-8"))
    assert len(evo["versions"]) == 3
    assert len(evo["threads"]) >= 2
    assert any("Sensibilidade CAPEX" in th["label"] for th in evo["threads"])
    print("OK structural diff/evolution tests passed")

if __name__ == "__main__":
    main()
